"""One-off script: seed signature.db from the existing ΠΥΘΑΓΟΡΑΣ ΣΤΑΥΡΟΣ.xlsx inventory sheet."""

from pathlib import Path

import openpyxl

from signature_app import db

XLSX_PATH = Path(__file__).resolve().parent.parent.parent / "ΠΥΘΑΓΟΡΑΣ ΣΤΑΥΡΟΣ.xlsx"


def main() -> None:
    db.init_db()

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["📦 ΑΠΟΘΗΚΗ"]

    imported = 0
    skipped = 0
    for code, category, house, original_name, signature_name, stock_ml, *_ in ws.iter_rows(
        min_row=3, max_row=ws.max_row, values_only=True
    ):
        if not code or not signature_name:
            skipped += 1
            continue
        if db.code_exists(code):
            skipped += 1
            continue
        db.add_product(
            code=code,
            category=category or "UNISEX",
            house=house,
            original_name=original_name,
            signature_name=signature_name,
            stock_ml=float(stock_ml) if stock_ml else 0.0,
            product_type=db.infer_product_type(code),
        )
        imported += 1

    print(f"Imported {imported} products, skipped {skipped} incomplete/duplicate rows.")

    import_transactions(wb, "ΑΓΟΡΕΣ", db.add_purchase, "purchases")
    import_transactions(wb, "🧾 ΠΩΛΗΣΕΙΣ", db.add_sale, "sales")


def import_transactions(wb, sheet_name: str, add_fn, label: str) -> None:
    """Import historical rows from a ΑΓΟΡΕΣ/ΠΩΛΗΣΕΙΣ-shaped sheet.

    ΑΡΧΙΚΟ STOCK in the inventory sheet is a starting baseline, not a live
    figure, so historical rows must adjust stock the same as new entries do.
    """
    ws = wb[sheet_name]

    imported = 0
    skipped = 0
    for _, date, code, _signature_name, _category, ml, comments in ws.iter_rows(
        min_row=4, max_row=ws.max_row, values_only=True
    ):
        if not date or not code or ml is None:
            continue
        if not db.code_exists(code):
            skipped += 1
            continue
        add_fn(
            date=date.date().isoformat(),
            code=code,
            ml=float(ml),
            comments=comments,
            adjust_stock=True,
        )
        imported += 1

    print(f"Imported {imported} {label}, skipped {skipped} rows with unknown codes.")


if __name__ == "__main__":
    main()
